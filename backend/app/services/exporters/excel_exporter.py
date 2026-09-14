"""
Service Layer: Excel Exporter (openpyxl)

Menghasilkan workbook spreadsheet Excel multi-tab profesional untuk Laporan Bulanan SiTernak:
- Tab 1: Ringkasan Eksekutif (Header, KPI Cards, Tabel Finansial dengan rumus SUM, Tabel Produksi Telur)
- Tab 2: Produksi & HDP Harian
- Tab 3: Rincian Penjualan
- Tab 4: Rincian Pengeluaran
Mengembalikan stream biner io.BytesIO in-memory.
"""

from io import BytesIO
from typing import Any, Dict
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Palette Warna & Format Standar
FONT_FAMILY = "Segoe UI"
COLOR_NAVY = "1E293B"       # Slate 800
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_MUTED_BG = "F8FAFC"   # Slate 50
COLOR_BORDER = "CBD5E1"     # Slate 300
COLOR_CARD_BORDER = "94A3B8"
COLOR_ACCENT = "0D9488"     # Teal 600

FORMAT_CURRENCY = '"Rp "#,##0'
FORMAT_INTEGER = '#,##0'
FORMAT_PERCENT = '0.0%'
FORMAT_DECIMAL = '#,##0.00'

BORDER_THIN = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="thin", color=COLOR_BORDER),
)

BORDER_TOTAL = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="double", color=COLOR_NAVY),
)

FILL_HEADER = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
FILL_ZEBRA = PatternFill(start_color=COLOR_MUTED_BG, end_color=COLOR_MUTED_BG, fill_type="solid")
FILL_CARD_BG = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

FONT_TITLE = Font(name=FONT_FAMILY, size=16, bold=True, color=COLOR_NAVY)
FONT_SUBTITLE = Font(name=FONT_FAMILY, size=11, bold=False, color="475569")
FONT_SECTION = Font(name=FONT_FAMILY, size=12, bold=True, color=COLOR_NAVY)
FONT_HEADER = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_HEADER_TEXT)
FONT_REGULAR = Font(name=FONT_FAMILY, size=10, bold=False, color="1E293B")
FONT_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="1E293B")
FONT_CARD_TITLE = Font(name=FONT_FAMILY, size=9, bold=True, color="64748B")
FONT_CARD_VALUE = Font(name=FONT_FAMILY, size=14, bold=True, color=COLOR_NAVY)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _autofit_columns(ws, min_width: int = 12, padding: int = 4):
    """
    Menyesuaikan lebar kolom secara otomatis berdasarkan isi sel terpanjang.
    """
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.number_format and ("Rp" in cell.number_format or "#,##0" in cell.number_format):
                # Tambahkan estimasi padding untuk angka yang diformat
                val_str += "    "
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + padding, min_width)


def generate_monthly_excel_report(data: Dict[str, Any]) -> BytesIO:
    """
    Menghasilkan workbook Excel multi-tab dari data laporan bulanan.
    """
    wb = openpyxl.Workbook()

    # =========================================================================
    # TAB 1: RINGKASAN EKSEKUTIF
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Ringkasan Eksekutif"
    ws1.views.sheetView[0].showGridLines = True

    # 1. Header Laporan
    ws1["A1"] = "SITERNAK - SISTEM MANAJEMEN PETERNAKAN AYAM PETELUR"
    ws1["A1"].font = FONT_TITLE
    ws1["A2"] = f"LAPORAN BULANAN KOMPREHENSIF - PERIODE: {data['periode_label'].upper()}"
    ws1["A2"].font = FONT_SUBTITLE
    ws1["A3"] = f"Tanggal Cetak: {data['tanggal_cetak']} | Standar Penilaian: SNI Layer Komersial"
    ws1["A3"].font = Font(name=FONT_FAMILY, size=9, italic=True, color="64748B")

    # 2. Blok KPI Cards (Baris 5-7)
    # Card 1: Laba Bersih MTD
    ws1["A5"] = "LABA / RUGI BERSIH"
    ws1["A5"].font = FONT_CARD_TITLE
    ws1["A6"] = data["finansial"]["laba_rugi_bersih"]
    ws1["A6"].font = FONT_CARD_VALUE
    ws1["A6"].number_format = FORMAT_CURRENCY
    ws1["A7"] = f"Margin: {data['finansial']['margin_persen']:.1f}% ({data['finansial']['status_pnl'].upper()})"
    ws1["A7"].font = Font(name=FONT_FAMILY, size=9, color="475569")

    # Card 2: Rata-Rata HDP
    ws1["C5"] = "RATA-RATA HDP BULANAN"
    ws1["C5"].font = FONT_CARD_TITLE
    ws1["C6"] = data["produksi"]["rata_rata_hdp"] / 100.0
    ws1["C6"].font = FONT_CARD_VALUE
    ws1["C6"].number_format = FORMAT_PERCENT
    ws1["C7"] = f"Total Panen: {data['produksi']['total_panen']:,} butir"
    ws1["C7"].font = Font(name=FONT_FAMILY, size=9, color="475569")

    # Card 3: FCR Pakan
    ws1["E5"] = "FEED CONVERSION RATIO (FCR)"
    ws1["E5"].font = FONT_CARD_TITLE
    ws1["E6"] = data["fcr"]["fcr"]
    ws1["E6"].font = FONT_CARD_VALUE
    ws1["E6"].number_format = FORMAT_DECIMAL
    ws1["E7"] = f"Status: {data['fcr']['status_efisiensi'].replace('_', ' ').title()} (Target: 2.10-2.35)"
    ws1["E7"].font = Font(name=FONT_FAMILY, size=9, color="475569")

    # Card 4: Saldo Akhir Stok Telur
    ws1["G5"] = "STOK AKHIR TELUR SIAP JUAL"
    ws1["G5"].font = FONT_CARD_TITLE
    ws1["G6"] = data["stok"]["saldo_akhir"]
    ws1["G6"].font = FONT_CARD_VALUE
    ws1["G6"].number_format = FORMAT_INTEGER
    ws1["G7"] = f"Status: {data['stok']['status'].upper()}"
    ws1["G7"].font = Font(name=FONT_FAMILY, size=9, color="475569")

    # Styling Card Blocks
    for col_start in ["A", "C", "E", "G"]:
        for r in range(5, 8):
            cell = ws1[f"{col_start}{r}"]
            cell.fill = FILL_CARD_BG
            cell.border = Border(
                left=Side(style="thin", color=COLOR_CARD_BORDER),
                right=Side(style="thin", color=COLOR_CARD_BORDER),
                top=Side(style="thin", color=COLOR_CARD_BORDER) if r == 5 else None,
                bottom=Side(style="thin", color=COLOR_CARD_BORDER) if r == 7 else None,
            )

    # 3. Tabel Ringkasan Finansial (Baris 9)
    ws1["A9"] = "1. RINGKASAN KEUANGAN BULANAN"
    ws1["A9"].font = FONT_SECTION

    fin_headers = ["Kategori Pos Keuangan", "Tipe Aliran", "Nominal (Rp)", "Rasio terhadap Pendapatan"]
    for i, h in enumerate(fin_headers, start=1):
        cell = ws1.cell(row=10, column=i, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER if i in [2, 4] else (ALIGN_RIGHT if i == 3 else ALIGN_LEFT)
        cell.border = BORDER_THIN

    row_idx = 11
    # Baris Pendapatan
    ws1.cell(row=row_idx, column=1, value="Pendapatan Penjualan Telur").font = FONT_BOLD
    ws1.cell(row=row_idx, column=2, value="Pemasukan").alignment = ALIGN_CENTER
    c_rev = ws1.cell(row=row_idx, column=3, value=data["finansial"]["total_pendapatan"])
    c_rev.number_format = FORMAT_CURRENCY
    c_rev.font = FONT_BOLD
    c_rev.alignment = ALIGN_RIGHT
    ws1.cell(row=row_idx, column=4, value=1.0).number_format = FORMAT_PERCENT
    ws1.cell(row=row_idx, column=4).alignment = ALIGN_CENTER
    for c in range(1, 5):
        ws1.cell(row=row_idx, column=c).border = BORDER_THIN

    # Baris Breakdown Pengeluaran
    start_exp_row = row_idx + 1
    for kat, nom in sorted(data["finansial"]["breakdown_kategori"].items()):
        row_idx += 1
        ws1.cell(row=row_idx, column=1, value=f"Beban {kat.replace('_', ' ').title()}").font = FONT_REGULAR
        ws1.cell(row=row_idx, column=2, value="Pengeluaran").alignment = ALIGN_CENTER
        c_exp = ws1.cell(row=row_idx, column=3, value=float(nom))
        c_exp.number_format = FORMAT_CURRENCY
        c_exp.font = FONT_REGULAR
        c_exp.alignment = ALIGN_RIGHT
        
        # Rasio terhadap pendapatan
        c_ratio = ws1.cell(row=row_idx, column=4, value=f"=C{row_idx}/C11" if data["finansial"]["total_pendapatan"] > 0 else 0.0)
        c_ratio.number_format = FORMAT_PERCENT
        c_ratio.alignment = ALIGN_CENTER
        for c in range(1, 5):
            ws1.cell(row=row_idx, column=c).border = BORDER_THIN

    end_exp_row = max(row_idx, start_exp_row)
    
    # Baris Total Pengeluaran (Formula Excel)
    row_idx += 1
    ws1.cell(row=row_idx, column=1, value="Total Beban Operasional").font = FONT_BOLD
    ws1.cell(row=row_idx, column=2, value="Subtotal").alignment = ALIGN_CENTER
    c_tot_exp = ws1.cell(row=row_idx, column=3, value=f"=SUM(C{start_exp_row}:C{end_exp_row})" if end_exp_row >= start_exp_row else data["finansial"]["total_pengeluaran"])
    c_tot_exp.number_format = FORMAT_CURRENCY
    c_tot_exp.font = FONT_BOLD
    c_tot_exp.alignment = ALIGN_RIGHT
    ws1.cell(row=row_idx, column=4, value=f"=C{row_idx}/C11" if data["finansial"]["total_pendapatan"] > 0 else 0.0).number_format = FORMAT_PERCENT
    ws1.cell(row=row_idx, column=4).alignment = ALIGN_CENTER
    for c in range(1, 5):
        ws1.cell(row=row_idx, column=c).border = BORDER_TOTAL

    # Baris Laba/Rugi Bersih
    row_idx += 1
    ws1.cell(row=row_idx, column=1, value="LABA / RUGI BERSIH BULANAN").font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY)
    ws1.cell(row=row_idx, column=2, value="Net Profit").alignment = ALIGN_CENTER
    c_net = ws1.cell(row=row_idx, column=3, value=f"=C11-C{row_idx-1}")
    c_net.number_format = FORMAT_CURRENCY
    c_net.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY)
    c_net.alignment = ALIGN_RIGHT
    ws1.cell(row=row_idx, column=4, value=f"=C{row_idx}/C11" if data["finansial"]["total_pendapatan"] > 0 else 0.0).number_format = FORMAT_PERCENT
    ws1.cell(row=row_idx, column=4).alignment = ALIGN_CENTER
    for c in range(1, 5):
        ws1.cell(row=row_idx, column=c).border = BORDER_TOTAL

    # 4. Tabel Ringkasan Produksi Telur
    row_idx += 3
    ws1.cell(row=row_idx, column=1, value="2. RINGKASAN PRODUKSI & MUTASI GUDANG").font = FONT_SECTION
    
    prod_headers = ["Indikator / Mutasi", "Kuantitas (Butir)", "Kuantitas (Kg Telur)", "Proporsi / Keterangan"]
    row_idx += 1
    for i, h in enumerate(prod_headers, start=1):
        cell = ws1.cell(row=row_idx, column=i, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER if i in [3, 4] else (ALIGN_RIGHT if i == 2 else ALIGN_LEFT)
        cell.border = BORDER_THIN

    tot_panen = max(data["produksi"]["total_panen"], 1)
    prod_rows = [
        ("Telur Normal (Grade A)", data["produksi"]["total_normal"], data["produksi"]["total_normal"] * 0.06, data["produksi"]["total_normal"] / tot_panen),
        ("Telur Retak (Grade B)", data["produksi"]["total_retak"], data["produksi"]["total_retak"] * 0.06, data["produksi"]["total_retak"] / tot_panen),
        ("Telur Pecah (Reject/Loss)", data["produksi"]["total_pecah"], data["produksi"]["total_pecah"] * 0.06, data["produksi"]["total_pecah"] / tot_panen),
        ("TOTAL PANEN FISIK", data["produksi"]["total_panen"], data["produksi"]["total_panen"] * 0.06, 1.0),
        ("Saldo Awal Gudang (Tgl 1)", data["stok"]["saldo_awal"], data["stok"]["saldo_awal"] * 0.06, "Stok Terbawa"),
        ("Penjualan Terdistribusi", data["stok"]["total_keluar_terjual"], data["stok"]["total_keluar_terjual"] * 0.06, "Pengurangan"),
        ("SALDO AKHIR GUDANG", data["stok"]["saldo_akhir"], data["stok"]["saldo_akhir"] * 0.06, f"Status: {data['stok']['status'].upper()}"),
    ]

    for label, butir, kg, prop in prod_rows:
        row_idx += 1
        is_bold = "TOTAL" in label or "SALDO AKHIR" in label
        c_lbl = ws1.cell(row=row_idx, column=1, value=label)
        c_lbl.font = FONT_BOLD if is_bold else FONT_REGULAR
        
        c_btr = ws1.cell(row=row_idx, column=2, value=butir)
        c_btr.font = FONT_BOLD if is_bold else FONT_REGULAR
        c_btr.number_format = FORMAT_INTEGER
        c_btr.alignment = ALIGN_RIGHT
        
        c_kg = ws1.cell(row=row_idx, column=3, value=kg)
        c_kg.font = FONT_BOLD if is_bold else FONT_REGULAR
        c_kg.number_format = FORMAT_DECIMAL
        c_kg.alignment = ALIGN_RIGHT
        
        c_prop = ws1.cell(row=row_idx, column=4, value=prop)
        c_prop.font = FONT_BOLD if is_bold else FONT_REGULAR
        c_prop.alignment = ALIGN_CENTER
        if isinstance(prop, float):
            c_prop.number_format = FORMAT_PERCENT

        for c in range(1, 5):
            ws1.cell(row=row_idx, column=c).border = BORDER_TOTAL if is_bold else BORDER_THIN

    _autofit_columns(ws1)

    # =========================================================================
    # TAB 2: PRODUKSI & HDP HARIAN
    # =========================================================================
    ws2 = wb.create_sheet(title="Produksi & HDP Harian")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = f"CATATAN PRODUKSI HARIAN & TREN HDP - {data['periode_label'].upper()}"
    ws2["A1"].font = FONT_TITLE
    ws2["A2"] = f"Rata-rata HDP: {data['produksi']['rata_rata_hdp']:.2f}% | Hari Tercatat: {data['produksi']['recorded_days']}/{data['produksi']['total_days']} Hari"
    ws2["A2"].font = FONT_SUBTITLE

    harian_headers = [
        "Tanggal", "Butir Normal", "Butir Retak", "Butir Pecah",
        "Total Panen", "Populasi Aktif", "HDP (%)", "Keterangan"
    ]
    for i, h in enumerate(harian_headers, start=1):
        cell = ws2.cell(row=4, column=i, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER if i in [1, 7, 8] else (ALIGN_RIGHT if i < 7 else ALIGN_LEFT)
        cell.border = BORDER_THIN

    for idx, d in enumerate(data["produksi"]["harian_produksi"], start=5):
        ws2.cell(row=idx, column=1, value=str(d["tanggal"])).alignment = ALIGN_CENTER
        ws2.cell(row=idx, column=2, value=d["normal"]).number_format = FORMAT_INTEGER
        ws2.cell(row=idx, column=3, value=d["retak"]).number_format = FORMAT_INTEGER
        ws2.cell(row=idx, column=4, value=d["pecah"]).number_format = FORMAT_INTEGER
        ws2.cell(row=idx, column=5, value=d["total"]).number_format = FORMAT_INTEGER
        ws2.cell(row=idx, column=6, value=d["populasi_aktif"]).number_format = FORMAT_INTEGER
        
        c_hdp = ws2.cell(row=idx, column=7, value=d["hdp_persen"] / 100.0)
        c_hdp.number_format = FORMAT_PERCENT
        c_hdp.alignment = ALIGN_CENTER
        
        ws2.cell(row=idx, column=8, value=d["keterangan"]).alignment = ALIGN_CENTER

        for c in range(1, 9):
            cell = ws2.cell(row=idx, column=c)
            cell.font = FONT_REGULAR
            cell.border = BORDER_THIN
            if idx % 2 == 0:
                cell.fill = FILL_ZEBRA

    _autofit_columns(ws2)

    # =========================================================================
    # TAB 3: RICIAN PENJUALAN
    # =========================================================================
    ws3 = wb.create_sheet(title="Rincian Penjualan")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = f"BUKU RINCIAN TRANSAKSI PENJUALAN - {data['periode_label'].upper()}"
    ws3["A1"].font = FONT_TITLE
    ws3["A2"] = f"Total Pendapatan: Rp {data['finansial']['total_pendapatan']:,.0f} | Total Butir Terdistribusi: {data['stok']['total_keluar_terjual']:,} butir"
    ws3["A2"].font = FONT_SUBTITLE

    sales_headers = [
        "Tanggal", "Pembeli", "Satuan", "Kuantitas",
        "Harga Satuan (Rp)", "Total (Rp)", "Butir Fisik Keluar"
    ]
    for i, h in enumerate(sales_headers, start=1):
        cell = ws3.cell(row=4, column=i, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER if i in [1, 3] else (ALIGN_RIGHT if i in [4, 5, 6, 7] else ALIGN_LEFT)
        cell.border = BORDER_THIN

    sales_start_row = 5
    cur_row = sales_start_row
    for s in data["detail_penjualan"]:
        ws3.cell(row=cur_row, column=1, value=str(s["tanggal"])).alignment = ALIGN_CENTER
        ws3.cell(row=cur_row, column=2, value=s["pembeli"]).alignment = ALIGN_LEFT
        ws3.cell(row=cur_row, column=3, value=s["satuan"].upper()).alignment = ALIGN_CENTER
        
        c_qty = ws3.cell(row=cur_row, column=4, value=s["kuantitas"])
        c_qty.number_format = FORMAT_DECIMAL if s["satuan"] == "kg" else FORMAT_INTEGER
        c_qty.alignment = ALIGN_RIGHT
        
        c_prc = ws3.cell(row=cur_row, column=5, value=s["harga_satuan"])
        c_prc.number_format = FORMAT_CURRENCY
        c_prc.alignment = ALIGN_RIGHT
        
        c_tot = ws3.cell(row=cur_row, column=6, value=s["total"])
        c_tot.number_format = FORMAT_CURRENCY
        c_tot.alignment = ALIGN_RIGHT
        
        c_btr = ws3.cell(row=cur_row, column=7, value=s["jumlah_butir"])
        c_btr.number_format = FORMAT_INTEGER
        c_btr.alignment = ALIGN_RIGHT

        for c in range(1, 8):
            cell = ws3.cell(row=cur_row, column=c)
            cell.font = FONT_REGULAR
            cell.border = BORDER_THIN
            if cur_row % 2 == 0:
                cell.fill = FILL_ZEBRA
        cur_row += 1

    # Baris Total Penjualan
    if cur_row > sales_start_row:
        ws3.cell(row=cur_row, column=1, value="TOTAL PENJUALAN").font = FONT_BOLD
        ws3.cell(row=cur_row, column=6, value=f"=SUM(F{sales_start_row}:F{cur_row-1})").number_format = FORMAT_CURRENCY
        ws3.cell(row=cur_row, column=6).font = FONT_BOLD
        ws3.cell(row=cur_row, column=6).alignment = ALIGN_RIGHT
        ws3.cell(row=cur_row, column=7, value=f"=SUM(G{sales_start_row}:G{cur_row-1})").number_format = FORMAT_INTEGER
        ws3.cell(row=cur_row, column=7).font = FONT_BOLD
        ws3.cell(row=cur_row, column=7).alignment = ALIGN_RIGHT
        for c in range(1, 8):
            ws3.cell(row=cur_row, column=c).border = BORDER_TOTAL

    _autofit_columns(ws3)

    # =========================================================================
    # TAB 4: RINCIAN PENGELUARAN
    # =========================================================================
    ws4 = wb.create_sheet(title="Rincian Pengeluaran")
    ws4.views.sheetView[0].showGridLines = True

    ws4["A1"] = f"BUKU RINCIAN BIAYA PENGELUARAN OPERASIONAL - {data['periode_label'].upper()}"
    ws4["A1"].font = FONT_TITLE
    ws4["A2"] = f"Total Beban: Rp {data['finansial']['total_pengeluaran']:,.0f} | Total Pakan Terpakai: {data['fcr']['total_kg_pakan']:,.1f} kg"
    ws4["A2"].font = FONT_SUBTITLE

    exp_headers = ["Tanggal", "Kategori", "Deskripsi", "Jumlah Kg (Pakan)", "Nominal Biaya (Rp)"]
    for i, h in enumerate(exp_headers, start=1):
        cell = ws4.cell(row=4, column=i, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER if i in [1, 2] else (ALIGN_RIGHT if i in [4, 5] else ALIGN_LEFT)
        cell.border = BORDER_THIN

    exp_start_row = 5
    cur_row = exp_start_row
    for e in data["detail_pengeluaran"]:
        ws4.cell(row=cur_row, column=1, value=str(e["tanggal"])).alignment = ALIGN_CENTER
        ws4.cell(row=cur_row, column=2, value=e["kategori"].replace("_", " ").title()).alignment = ALIGN_CENTER
        ws4.cell(row=cur_row, column=3, value=e["deskripsi"]).alignment = ALIGN_LEFT
        
        c_kg = ws4.cell(row=cur_row, column=4, value=e["jumlah_kg"] if e["jumlah_kg"] is not None else "-")
        if isinstance(e["jumlah_kg"], (int, float)):
            c_kg.number_format = FORMAT_DECIMAL
            c_kg.alignment = ALIGN_RIGHT
        else:
            c_kg.alignment = ALIGN_CENTER
            
        c_nom = ws4.cell(row=cur_row, column=5, value=e["nominal"])
        c_nom.number_format = FORMAT_CURRENCY
        c_nom.alignment = ALIGN_RIGHT

        for c in range(1, 6):
            cell = ws4.cell(row=cur_row, column=c)
            cell.font = FONT_REGULAR
            cell.border = BORDER_THIN
            if cur_row % 2 == 0:
                cell.fill = FILL_ZEBRA
        cur_row += 1

    # Baris Total Pengeluaran
    if cur_row > exp_start_row:
        ws4.cell(row=cur_row, column=1, value="TOTAL PENGELUARAN").font = FONT_BOLD
        ws4.cell(row=cur_row, column=5, value=f"=SUM(E{exp_start_row}:E{cur_row-1})").number_format = FORMAT_CURRENCY
        ws4.cell(row=cur_row, column=5).font = FONT_BOLD
        ws4.cell(row=cur_row, column=5).alignment = ALIGN_RIGHT
        for c in range(1, 6):
            ws4.cell(row=cur_row, column=c).border = BORDER_TOTAL

    _autofit_columns(ws4)

    # Simpan ke stream BytesIO
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
